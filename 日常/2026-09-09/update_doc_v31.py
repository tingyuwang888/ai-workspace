#!/usr/bin/env python3
"""将 天策系统策略测试组件覆盖能力说明.docx 从 V3.0 升级为 V3.1（组件族全景 + 治理五态 + 两张工作流图）"""
import copy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = "/Users/td/Desktop/天策系统策略测试组件覆盖能力说明.docx"
OUT = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/outputs/天策系统策略测试组件覆盖能力说明_V3.1.docx"
FIG1 = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/tiance_metric_policy_flow.png"
FIG2 = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/tiance_component_family.png"

doc = Document(SRC)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise AssertionError(f"anchor not found: {prefix}")

def set_text(p, text):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(text)

def insert_para_after(anchor, text="", style=None, center=False):
    new_p = OxmlElement('w:p')
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if text:
        para.add_run(text)
    if style:
        para.style = doc.styles[style]
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para

def chain(anchor, items, style=None, center=False):
    """items: list of str; insert sequentially after anchor, return last"""
    last = anchor
    for t in items:
        last = insert_para_after(last, t, style=style, center=center)
    return last

def insert_table_after(anchor, data):
    rows, cols = len(data), len(data[0])
    tbl = doc.add_table(rows=rows, cols=cols)
    anchor._p.addnext(tbl._tbl)
    try:
        tbl.style = doc.tables[1].style
    except Exception:
        pass
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            tbl.cell(ri, ci).text = val
    # header shading + bold white
    for ci in range(cols):
        cell = tbl.cell(0, ci)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), '1A5276')
        tcPr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)
    # body font size 9
    for ri in range(1, rows):
        for ci in range(cols):
            for p in tbl.cell(ri, ci).paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return tbl

def insert_picture_after(anchor, path, width=6.3):
    doc.add_picture(path, width=Inches(width))
    pic_para = doc.paragraphs[-1]
    anchor._p.addnext(pic_para._p)
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return pic_para

# ---------- 表 0 元数据 ----------
t0 = doc.tables[0]
t0.cell(1, 0).text = "V3.1"
t0.cell(1, 2).text = "2026年9月9日"

# ---------- §一 概览 ----------
set_text(find("目前，策略测试能力"),
    "目前，天策测试组件族已覆盖从组件生成（字段/数据源/ETL/函数/指标/规则/策略/接口服务）到测试执行（用例生成/平台执行/报告审计/闭环编排）的完整链路，形成\u201c生成即可测、测完可追溯\u201d的治理闭环。")
set_text(find("完整链路"),
    "完整链路  元数据核对 → 组件生成（8类）→ 平台版本核对（policyCode/policyVersion/bizType，快照≤10分钟）→ 用例设计与覆盖门禁 → Fixture 隔离准备 → 逐笔执行与证据采集 → 五态复核与根因分析 → 报告质量审计 → 完成门禁 → 最小重跑（caseKey 精确定位）")

# ---------- §二.1 字段 ----------
chain(find("业务示例：输入收入和负债"),
      ["对应组件  tiance-field-forge（系统字段 + 动态字段 + 对象字段 + 枚举字段）。"])

# ---------- §二.2 外部数据 ----------
set_text(find("当前边界"),
    "当前边界  默认 no-mock 模式：不再生成三方超时/错误码/字段类型错误用例；不可控三方候选进入 excludedScenarios 并计入覆盖缺口，禁止通过过滤候选提高覆盖率。真实接口的超时、网络故障和错误码，需要专用模拟或隔离环境。")
chain(find("当前边界"),
      ["对应组件  tiance-datasource-forge（.ds 生成 + DES 加密 + JSON-ETL 协作 + tar 打包）。"])

# ---------- §二.3 数据处理 ----------
chain(find("业务示例：查询多笔司法案件"),
      ["对应组件  tiance-etl-forge（前置/后置/分页 ETL + Java/Groovy 代码生成 + outputMap/serviceParam 对齐 + 反向验证）。"])

# ---------- §二.4 函数 ----------
chain(find("验证多笔数据汇总"),
      ["多行场景约束  仅当模板明确声明函数入参为数组/列表/集合，或同一查询键返回多行记录时才生成多行场景；SUM/汇总/合计/分组不能单独证明函数支持多行输入。"])
chain(find("判断原则"),
      ["对应组件  tiance-function-forge（.fun 交付 + 权威 Excel 唯一业务源 + 函数测试工作簿）。"])

# ---------- §二.5 实时指标（8 bullets + 边界 + 对应组件 + 图1） ----------
p45 = find("验证统计对象")
p46 = find("验证命中、不命中")
p47 = find("验证实际指标值")
set_text(p45, "验证统计对象、时间范围、计算字段、筛选条件。")
set_text(p46, "验证当前笔计入方式（含/不含/条件含）。")
set_text(p47, "验证历史贡献与最终算术分开列（历史值 + 当前笔 = 最终值）。")
last = p47
last = chain(last, [
    "验证事件过滤、窗口与排除窗口。",
    "验证去重维度、聚合键、名单/标签、动态阈值。",
    "验证时区与窗口边界。",
    "验证命中、不命中、精确阈值、阈值前后值。",
    "隔离全部指标分组维度，不能只换业务 ID。",
])
last = chain(last, [
    "实现方式与边界  指标验证嵌在策略执行链路中（通过目标规则命中反推指标值），不覆盖独立 .zb 回归、指标计算 API 直测、性能/并发/TPS。",
    "对应组件  tiance-realtime-metric-forge（.zb 生成，10 类模板路由）。",
])
pic1 = insert_picture_after(last, FIG1)
chain(pic1, ["图 1  带指标的策略测试执行链路（指标值三段式验证）"], center=True)

# ---------- §二.6 规则和规则集 ----------
last = chain(find("验证伴随规则"), [
    "三方节点证据使用 fieldValues / absentFields 精确断言，Excel \u201c三方数据\u201d列已改名为\u201c预期执行证据\u201d。",
    "覆盖率按规则集+规则编码统计，不跨规则集合并；applicableTags 与 notApplicableTags 分开报告，不适用场景不计为缺失。",
])
chain(find("业务示例：年龄小于18岁"),
      ["对应组件  tiance-rule-forge（.rss 生成 + autoCreateMetrics 联动）。"])

# ---------- §二.7 决策流 ----------
chain(find("验证流程断点"), [
    "集中度策略专用语义场景  行政区全称/简称兼容、特殊区域 OR 分支、多行聚合、聚合函数小数精度、仅实时/仅离线数据；判定适用的场景必须使用真实 MySQL Fixture。",
])
chain(find("业务示例：先查询工商"),
      ["对应组件  tiance-policy-forge（.pls 简易/流模式 + ProcessOn .pos/.posf/VSDX 转换 + 循环子策略）。"])

# ---------- §二.8 接口服务 ----------
chain(find("业务示例：业务系统传入"),
      ["对应组件  tiance-service-config-forge（serviceConfig CSV + ServiceFieldMapping + creditEngine/transEngine）。"])

# ---------- 章节重编号（先做，避免与新§三标题冲突） ----------
set_text(find("三、当前已经形成的测试闭环"), "四、当前已经形成的测试闭环")
set_text(find("四、正确理解的边界"), "五、正确理解的边界")
set_text(find("五、汇报结论"), "六、汇报结论")

# ---------- 原§三（现§四）补映射句 ----------
chain(find("生成报告，分析失败原因"),
      ["上述业务视角九步在技术层对应§三\u201c治理闭环 8 阶段\u201d表，两者一一映射。"])

# ---------- 原§四（现§五）停止条件 + 不覆盖清单 ----------
set_text(find("停止条件"),
    "停止条件  如果渠道、机构、字段、外部服务、名单、指标样本或其他依赖尚未准备好，测试会明确列出缺口并停止，不会使用猜测数据继续执行。平台快照超 10 分钟阻断；Fixture run ID 每轮新建，禁止复用上一轮隔离键。")
last = chain(find("真实第三方异常需要模拟环境"), ["明确不覆盖清单"])
insert_table_after(last, [
    ["不覆盖项", "原因", "替代方案"],
    ["独立 .zb 指标回归", "无独立指标执行入口，须绑策略", "新建 tiance-metric-standalone-test skill"],
    ["指标计算 API 直测", "现有链路只走 /noah/policyTest", "摸 salaxy/indexApi 契约后新建"],
    ["指标性能/并发/TPS", "不在功能测试范围", "JMeter/K6 独立方案"],
    ["模型效果评估", "需 A/B 与样本回溯", "独立建模评估流程"],
    ["真实第三方异常（超时/网络故障/错误码）", "默认 no-mock，不生成异常用例", "专用模拟或隔离环境"],
    ["生产策略/名单/路由修改、存量数据清理", "不属于测试授权范围", "走变更管理流程"],
])

# ---------- 原§五（现§六）汇报结论 + 表3 五态 ----------
set_text(find("业务化概括"),
    "业务化概括  业务人员说明在哪个渠道、针对什么对象、需要哪些数据、按照什么条件判断、最终期望返回什么结果，系统可以协助形成测试场景、准备测试数据、执行平台验证，并输出一套可解释、可追溯的测试结论。每条用例保留稳定 caseKey、设计清单、执行证据与治理状态；最小重跑基于 caseKey 精确修改，不按关键词或同规则第一条用例定位。")
set_text(find("最终判断原则"),
    "最终判断原则  只有目标指标、目标规则和最终策略结果均有证据并符合预期，测试才判定通过。（治理合同五态定义详见§六结果分类表）")

def tbl_by_header(h0):
    for t in doc.tables:
        if t.rows[0].cells[0].text.strip() == h0:
            return t
    raise AssertionError(f"table not found by header: {h0}")

t3 = tbl_by_header("结果")
new3 = [
    ["通过", "目标指标、规则和策略结果均有证据且符合预期", "纳入通过统计"],
    ["失败", "实际结果与业务预期存在已证实差异", "根因分析（数据/配置/平台/机制），产出 analysis.json"],
    ["执行阻塞", "环境/接口/数据库/平台故障或 401/超时导致本条未完成", "修复环境后按 caseKey 最小重跑"],
    ["编排阻塞", "设计门禁/平台门禁/Fixture 门禁未通过，用例未提交", "补齐前置条件后进入 B/C/D 阶段"],
    ["无效用例", "设计歧义、caseKey 冲突、覆盖重复或场景不适用", "记入 excludedScenarios 并计入覆盖缺口"],
]
for ri, row in enumerate(new3, start=1):
    for ci, val in enumerate(row):
        t3.cell(ri, ci).text = val

# ---------- 新§三 组件族与治理闭环（插在 现§四 标题之前） ----------
anchor_h4 = find("四、当前已经形成的测试闭环")
def insert_para_before(anchor, text="", style=None, center=False):
    new_p = OxmlElement('w:p')
    anchor._p.addprevious(new_p)
    para = Paragraph(new_p, anchor._parent)
    if text:
        para.add_run(text)
    if style:
        para.style = doc.styles[style]
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para

def insert_table_before(anchor, data):
    rows, cols = len(data), len(data[0])
    tbl = doc.add_table(rows=rows, cols=cols)
    anchor._p.addprevious(tbl._tbl)
    try:
        tbl.style = doc.tables[1].style
    except Exception:
        pass
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            tbl.cell(ri, ci).text = val
    for ci in range(cols):
        cell = tbl.cell(0, ci)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'1A5276')
        tcPr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); r.font.size = Pt(9)
    for ri in range(1, rows):
        for ci in range(cols):
            for p in tbl.cell(ri, ci).paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return tbl

def insert_picture_before(anchor, path, width=6.3):
    doc.add_picture(path, width=Inches(width))
    pic_para = doc.paragraphs[-1]
    anchor._p.addprevious(pic_para._p)
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return pic_para

insert_para_before(anchor_h4, "三、组件族与治理闭环", style="Heading 1")
insert_para_before(anchor_h4,
    "天策测试组件族由 14 个 skill 构成，分为支撑层（元数据与生命周期）、生成层（8 类组件 forge）和执行层（4 个 metric-* 测试 skill）。所有 skill 遵循同一份治理合同，产出物通过稳定 caseKey 串联。")
insert_table_before(anchor_h4, [
    ["层级", "Skill 名", "产出物", "在测试链路中的角色"],
    ["支撑层", "tiance-metadata-scout", "metadata_latest JSON 缓存", "提供字段/规则/规则集/策略/函数/数据源/合作方/合同/外数指标元数据"],
    ["支撑层", "tiance-component-lifecycle", "导入/上线/下线/删除操作", "组件全生命周期总调度"],
    ["生成层", "tiance-field-forge", "field_mapping.json + 系统字段.xls + 动态字段.xls", "字段定义"],
    ["生成层", "tiance-datasource-forge", ".ds 文件 + tar 包", "外部数据接口"],
    ["生成层", "tiance-etl-forge", ".etl 文件（前置/后置/分页）", "数据加工"],
    ["生成层", "tiance-function-forge", ".fun 文件 + 函数测试工作簿", "可复用计算逻辑"],
    ["生成层", "tiance-realtime-metric-forge", ".zb 文件 + 指标配置 Excel", "实时指标（10 类模板）"],
    ["生成层", "tiance-rule-forge", ".rss 文件 + rule_rss 产物", "规则集"],
    ["生成层", "tiance-policy-forge", ".pls 文件（简易/流模式）", "决策流策略"],
    ["生成层", "tiance-service-config-forge", "serviceConfig CSV", "接口服务配置"],
    ["执行层", "tiance-metric-testcase-generator", "testcases.json/xlsx + design-manifest.json + coverage-plan.json", "用例设计与覆盖门禁"],
    ["执行层", "tiance-metric-policy-test", "executionManifest + Fixture + 15 列 Excel 报告", "平台执行与证据采集"],
    ["执行层", "tiance-metric-report-checker", "auditManifest + 质量检查 Sheet", "报告审计（4 维度）"],
    ["执行层", "tiance-metric-agent-loop", "convergence.json + rerun-plan.json", "治理闭环编排（8 阶段）"],
])
insert_picture_before(anchor_h4, FIG2)
insert_para_before(anchor_h4, "图 2  天策测试组件族工作流（生成→执行→治理闭环）", center=True)
insert_para_before(anchor_h4, "治理闭环由 tiance-metric-agent-loop 编排，8 阶段串行 + 反馈回环：")
insert_table_before(anchor_h4, [
    ["阶段", "名称", "责任 skill", "输出产物", "硬门禁"],
    ["A", "范围与备份", "agent-loop", "备份目录 + coverage-plan.json", "原 Excel 与运行归档带时间戳备份，新结果写新 run 目录"],
    ["B", "设计/覆盖门禁", "testcase-generator", "design-manifest.json + design-review.json + testcases.governed.json", "每条规则 1 正 1 反为最低基线；边界案例显式标正反极性；设计待确认不提交"],
    ["C", "平台/执行门禁", "policy-test", "platform_snapshot + precheck_result", "policyCode/policyVersion/bizType/发布状态核对，快照 ≤10 分钟；名称缺失阻断"],
    ["D", "Fixture 准备", "policy-test", "fixtureManifest + 新 fixtureRunId", "首轮人工确认库连接/隔离键/清理范围；禁止复用上一轮隔离键"],
    ["E", "逐笔执行", "policy-test", "executionManifest + segments", "每 10 条分段；401 停；--max-retry 0；--design-manifest 必传"],
    ["F", "五态复核与证据分析", "report-checker", "governed-results.json + analysis.json", "只使用通过/失败/执行阻塞/编排阻塞/无效用例；未提交计划内用例不能从统计消失"],
    ["G", "报告质量审计", "report-checker", "auditManifest + 质量检查 Sheet", "4 维度（判定一致性/数据完整性/语义重复/实际结果质量）；安全违约非空即停"],
    ["H", "完成门禁与最小重跑", "agent-loop", "convergence.json + rerun-plan.json + cleanupResultPath", "validate-rerun 通过；caseKey 精确定位；未清理不得 completed"],
])

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", OUT)
