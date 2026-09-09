#!/usr/bin/env python3
"""V3.4：镜像《天策自动化配置能力概览_提交Chen》模板重排测试能力说明（散文为主、2 表、按Skill支持/边界）"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/outputs/天策系统策略测试组件覆盖能力说明_V3.4.docx"
FIG1 = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/tiance_metric_policy_flow.png"
FIG2 = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/tiance_component_family.png"

doc = Document()
st = doc.styles['Normal']
st.font.name = '微软雅黑'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def para(text="", style=None, center=False):
    p = doc.add_paragraph(style=style)
    if text: p.add_run(text)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def caption(text):
    try:
        return para(text, style='Caption', center=True)
    except KeyError:
        p = para(text, center=True)
        for r in p.runs: r.font.size = Pt(9); r.font.italic = True
        return p

def add_table(data):
    rows, cols = len(data), len(data[0])
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    for ri, row in enumerate(data):
        for ci, v in enumerate(row):
            t.cell(ri, ci).text = v
    for ci in range(cols):
        cell = t.cell(0, ci)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'1A5276')
        tcPr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold=True; r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); r.font.size=Pt(9)
    for ri in range(1, rows):
        for ci in range(cols):
            for p in t.cell(ri, ci).paragraphs:
                for r in p.runs: r.font.size = Pt(9)
    return t

def add_fig(path, cap):
    doc.add_picture(path, width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(cap)

# ===== 标题与定位 =====
para("天策策略测试能力概览", style='Title', center=True)
para("提交 Chen    2026年9月9日", center=True)
para("我们已形成覆盖字段、外部数据、数据处理、函数、实时指标、规则与规则集、决策流和接口服务的策略测试能力，可按治理闭环完成用例设计、隔离执行、证据采集与报告审计。Skill 是这些可复用测试工具的载体；业务人员确认预期后，工具协助生成用例、执行平台验证并按五态输出可追溯结论。")
para("当前定位是有人员确认的测试与交付辅助能力。它可以减少重复测试与证据整理，但不替代业务口径确认、真实环境验证和上线审批。")

# ===== 一 =====
doc.add_heading("一 现有测试能力如何覆盖业务决策", level=1)
add_fig(FIG1, "图1 带指标的策略测试执行链路")
add_table([
    ["测试对象", "现在可以验证什么", "典型业务示例"],
    ["字段", "普通/枚举/对象/动态字段的正常值、空值、非法值与边界值；动态字段计算结果。", "输入收入与负债，验证负债率并覆盖空值、零值与阈值。"],
    ["外部数据", "请求参数映射、返回码/字段、空/单/多结果、分页与结果转换，及进入函数规则后的业务影响。", "按统一社会信用代码查企业信息，验证状态与注册资本被准入规则正确使用。"],
    ["数据处理", "请求前转换、返回后提取与类型转换、多页累计与明细合并的清洗结果。", "汇总多页未结案司法案件金额，验证规则使用值与人工计算一致。"],
    ["函数", "函数节点实际执行、输入字段、输出值与允许误差、空值/非法值与多笔汇总。", "验证收入负债比、多笔交易金额合计、身份证年龄与距到期日天数。"],
    ["实时指标", "统计对象/时间范围/计算字段/筛选条件、当前笔计入、历史贡献与最终算术、窗口/去重/时区/分组隔离。", "验证近30天申请次数的历史贡献+当前笔=最终值。"],
    ["规则与规则集", "普通条件、多条件组合、名单判断、分支/终止/顺序/权重，正反例命中目标规则与伴随/禁止命中。", "年龄小于18拒绝；命中黑名单且金额超10万进入高风险。"],
    ["决策流", "固定顺序与条件分支流程、函数/规则集/外部数据/子流程调用、循环明细、断点与回环。", "先查工商司法，再算近30天次数与负债率，执行准入后分流通过/拒绝/人工。"],
    ["接口服务", "服务入参与平台字段映射、必填/固定值/空值、输出结果/风险等级/命中原因。", "业务系统传客户编号与金额，平台返回通过/拒绝、等级与原因。"],
])
para("上述为测试能力覆盖范围，不代表每类组件的所有组合都已在全部平台版本和项目中验证。")

# ===== 二 =====
doc.add_heading("二 测试边界与正式使用条件", level=1)
para("在逐笔执行之外，我们已有配套工具完成用例设计、覆盖门禁、Fixture 隔离与报告审计，并按治理合同输出五态结论。机构、渠道、产品等基础配置与组件生成属于支撑能力，不计入本文测试范围。")
add_fig(FIG2, "图2 测试治理闭环（执行 → A-H 八阶段）")
doc.add_heading("可以自动化的环节", level=2)
para("在资料明确、依赖齐备且场景属于支持范围时，可协助解析落地方案、生成测试用例、准备隔离 Fixture、逐笔提交并采集证据、比对目标节点入参与指标值、输出带证据的测试报告。人员仍需确认判断阈值、计算口径、隔离键与清理范围。")
doc.add_heading("指标测试并非所有模板完全一致", level=2)
para("函数、规则与规则集、决策流已具备较完整的证据比对能力。实时指标按三段式（历史贡献+当前笔贡献=最终算术值）验证并隔离全部分组维度，但验证嵌在策略执行链路中；独立 .zb 回归、指标计算 API 直测与性能/并发目前不支持，其他指标模板须按具体场景逐项确认。")
doc.add_heading("依赖缺失或场景超出范围时需要补齐", level=2)
para("字段、外部服务、名单、指标样本和运行组件必须真实可用。平台快照超 10 分钟阻断；Fixture run ID 每轮新建，禁止复用上一轮隔离键。默认 no-mock：不生成三方超时/错误码/字段类型错误用例，不可控三方计入覆盖缺口而非过滤掉。")
doc.add_heading("执行完成不等于正式上线", level=2)
para("正式使用仍需在目标平台核对业务预期，经业务负责人和平台负责人验收，发布后再次回查接口与策略结果。不同版本、项目和外部接口的适配需分别确认，不能把一次验证推广为全部兼容，也不能把工具执行权限视为上线审批。")

# ===== 三 =====
doc.add_heading("三 具体能力说明", level=1)
para("下面按 Skill 独立说明。每节先说明当前支持范围，再列出明确不支持或仍需外部确认的内容。")
doc.add_heading("用例生成 Skill  tiance-metric-testcase-generator", level=2)
para("支持：从落地方案 Excel 生成规则级、函数级、决策流分支、预警等级与指标级用例，输出 testcases 与 design-manifest、coverage-plan；每条规则 1 正 1 反为最低基线，边界案例显式标正反极性。")
para("边界：不自动授权平台提交或 Fixture  setup/cleanup；设计待确认不提交；不可控三方进入 excludedScenarios 并计入覆盖缺口。")
doc.add_heading("平台执行 Skill  tiance-metric-policy-test", level=2)
para("支持：平台门禁核对 policyCode/policyVersion/bizType/发布状态；按 design-manifest 准备隔离 Fixture；逐笔提交 /noah/policyTest 并采集目标节点入参、指标值与命中证据；每 10 条分段、401 停、--max-retry 0。")
para("边界：不修改生产策略/名单/路由，不清理存量数据；证据查询不重新提交用例；结果损坏或提交不确定时阻断而非自动重试。")
doc.add_heading("报告审计 Skill  tiance-metric-report-checker", level=2)
para("支持：对报告执行判定一致性、数据完整性、语义重复、实际结果质量四维审计；逐层比对设计→生成→实际请求→节点入参/指标；历史指标额外核对当前笔计入、分组去重键、窗口时区与过滤条件。")
para("边界：关键词反馈不能充当平台缺陷证据；日志未取到记为证据阻塞而非未执行；默认仅展示经复核的严重问题。")
doc.add_heading("闭环编排 Skill  tiance-metric-agent-loop", level=2)
para("支持：编排 A-H 八阶段（范围备份→设计门禁→平台门禁→Fixture→逐笔执行→五态复核→报告审计→完成重跑），产出 governed-results、analysis、convergence、rerun-plan；失败按 caseKey 最小重跑。")
para("边界：不按报告问题数或通过率自动宣告完成；adjustExpected 必须有规则定义依据，禁止按实际值反推预期；不明机制只安排经授权的单变量探针。")

# ===== 四 =====
doc.add_heading("四 使用边界和完成判断", level=1)
para("测试链路的前半段是用例设计、Fixture 准备与本地预检，后半段是目标平台的逐笔执行、证据采集、五态复核与报告审计。两段之间需要人员确认和平台权限，不能由生成结果自动推断。")
para("总体判断：现有能力已经覆盖从数据准备到决策结果验证的主要工作，可以显著减少重复测试与证据整理工作。正式使用的边界仍然清晰：业务口径由人员确认，平台状态以实际执行、证据与回查为准，未验证的组合能力不作承诺。")
add_table([
    ["状态", "可以说明什么", "不能说明什么"],
    ["通过", "目标指标、规则和策略结果均有证据且符合预期。", "不代表其他版本、项目或外部接口自动兼容。"],
    ["失败", "实际结果与业务预期存在已证实差异，已定位根因。", "不代表平台缺陷，需区分数据/配置/平台/机制。"],
    ["执行阻塞", "环境、接口、数据库或平台故障导致本条未完成。", "不代表目标规则未执行；证据缺失记阻塞不记未通过。"],
    ["编排阻塞", "设计/平台/Fixture 门禁未通过，用例未提交。", "不能用猜测值或历史文件替代前置确认。"],
    ["无效用例", "设计歧义、caseKey 冲突、覆盖重复或场景不适用。", "不能从统计中消失，须记入覆盖缺口。"],
])
para("参考资料  当前能力概览；天策策略测试组件覆盖能力说明 V3.3；tiance-metric-* 四 Skill 治理合同（governance.py / risk-governance.md）。")

doc.save(OUT)
print("saved:", OUT)
