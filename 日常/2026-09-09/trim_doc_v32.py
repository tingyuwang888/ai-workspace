#!/usr/bin/env python3
"""V3.1 → V3.2：只保留测试相关内容。移除支撑层/生成层（组件表行、对应组件行、图2泳道），精简概览。"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/outputs/天策系统策略测试组件覆盖能力说明_V3.1.docx"
OUT = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/outputs/天策系统策略测试组件覆盖能力说明_V3.2.docx"
FIG2 = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/tiance_component_family.png"

doc = Document(SRC)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise AssertionError("anchor not found: " + prefix)

def set_text(p, text):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(text)

def remove_para(p):
    p._p.getparent().remove(p._p)

def tbl_by_header(h0):
    for t in doc.tables:
        if t.rows[0].cells[0].text.strip() == h0:
            return t
    return None

def style_table(tbl):
    try:
        tbl.style = doc.tables[1].style
    except Exception:
        pass
    cols = len(tbl.columns)
    for ci in range(cols):
        cell = tbl.cell(0, ci)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'1A5276')
        tcPr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); r.font.size = Pt(9)
    for ri in range(1, len(tbl.rows)):
        for ci in range(cols):
            for p in tbl.cell(ri, ci).paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

# 1) 版本
doc.tables[0].cell(1, 0).text = "V3.2"

# 2) 概览精简
set_text(find("目前，天策测试组件族已覆盖"),
    "目前，天策策略测试能力已覆盖从用例设计、平台执行、证据采集到报告审计与闭环编排的完整链路，形成\u201c测完可追溯\u201d的治理闭环。")
set_text(find("完整链路"),
    "完整链路  平台版本核对（policyCode/policyVersion/bizType，快照≤10分钟）→ 用例设计与覆盖门禁 → Fixture 隔离准备 → 逐笔执行与证据采集 → 五态复核与根因分析 → 报告质量审计 → 完成门禁 → 最小重跑（caseKey 精确定位）")

# 3) 删除所有 对应组件 行（生成层引用）
for p in list(doc.paragraphs):
    if p.text.strip().startswith("对应组件"):
        remove_para(p)

# 4) §三 引言精简
set_text(find("天策测试组件族由"),
    "天策策略测试由 4 个测试执行 skill 构成，遵循同一份治理合同，产出物通过稳定 caseKey 串联。")

# 5) 重建组件表：仅执行层 4 行，3 列
old = tbl_by_header("层级")
assert old is not None
anchor_el = old._tbl.getprevious()
anchor = Paragraph(anchor_el, old._tbl.getparent()) if anchor_el.tag.endswith('}p') else None
# 若前一元素不是段落（不太可能），回退到引言段落
if anchor is None:
    anchor = find("天策策略测试由")
old._tbl.getparent().remove(old._tbl)
data = [
    ["Skill 名", "产出物", "在测试链路中的角色"],
    ["tiance-metric-testcase-generator", "testcases.json/xlsx + design-manifest.json + coverage-plan.json", "用例设计与覆盖门禁"],
    ["tiance-metric-policy-test", "executionManifest + Fixture + 15 列 Excel 报告", "平台执行与证据采集"],
    ["tiance-metric-report-checker", "auditManifest + 质量检查 Sheet", "报告审计（4 维度）"],
    ["tiance-metric-agent-loop", "convergence.json + rerun-plan.json", "治理闭环编排（8 阶段）"],
]
tbl = doc.add_table(rows=len(data), cols=3)
anchor._p.addnext(tbl._tbl)
for ri, row in enumerate(data):
    for ci, v in enumerate(row):
        tbl.cell(ri, ci).text = v
style_table(tbl)

# 6) 替换图 2 图片与图注
cap = find("图 2")
pic_el = cap._p.getprevious()
cap._p.getparent().remove(pic_el)
doc.add_picture(FIG2, width=Inches(6.3))
pic_para = doc.paragraphs[-1]
cap._p.addprevious(pic_para._p)
pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_text(cap, "图 2  天策测试工作流（执行 → 治理闭环）")

doc.save(OUT)
print("saved:", OUT)
