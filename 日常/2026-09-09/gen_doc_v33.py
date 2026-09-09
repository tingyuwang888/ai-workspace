#!/usr/bin/env python3
"""V3.3 重排：§二 压缩为"两个工作流覆盖组件"矩阵+双图；删除八个小节/9步闭环/指标模板表；五章精简版"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/outputs/天策系统策略测试组件覆盖能力说明_V3.3.docx"
FIG1 = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/tiance_metric_policy_flow.png"
FIG2 = "/Users/td/.qoderwork/workspace/mttrqawk3qp6mzcs/tiance_component_family.png"

doc = Document()

# 全局字体
st = doc.styles['Normal']
st.font.name = '微软雅黑'
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def para(text="", style=None, center=False, bold=False, size=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        if bold: r.font.bold = True
        if size: r.font.size = Pt(size)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def add_table(data, widths=None):
    rows, cols = len(data), len(data[0])
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    for ri, row in enumerate(data):
        for ci, v in enumerate(row):
            t.cell(ri, ci).text = v
    for ci in range(cols):
        cell = t.cell(0, ci)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'1A5276')
        tcPr.append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); r.font.size = Pt(9)
    for ri in range(1, rows):
        for ci in range(cols):
            for p in t.cell(ri, ci).paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return t

def add_fig(path, caption):
    doc.add_picture(path, width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(caption, center=True, size=9)

# ===== 标题 + 元数据 =====
para("策略测试能力概览", style='Title', center=True)
para("从数据准备到决策结果验证（测试执行视角）", center=True, size=11)
add_table([
    ["文档版本", "资料属性", "更新时间"],
    ["V3.3", "业务交流版", "2026年9月9日"],
])

# ===== 一、概览 =====
doc.add_heading("一、当前测试能力概览", level=1)
para("天策策略测试能力已覆盖从用例设计、平台执行、证据采集到报告审计与闭环编排的完整链路，形成\u201c测完可追溯\u201d的治理闭环。")
para("完整链路  平台版本核对（policyCode/policyVersion/bizType，快照≤10分钟）→ 用例设计与覆盖门禁 → Fixture 隔离准备 → 逐笔执行与证据采集 → 五态复核与根因分析 → 报告质量审计 → 完成门禁 → 最小重跑（caseKey 精确定位）")
add_table([
    ["阶段", "测试关注点"],
    ["数据准备", "测试字段是否完整，历史数据和外部数据是否正确进入平台"],
    ["计算加工", "ETL、指标和函数是否得到符合业务预期的结果"],
    ["风险判断", "规则、规则集是否按配置执行"],
    ["流程编排", "不同客户和业务场景是否走到正确决策分支"],
    ["结果交付", "最终结果、命中原因和测试证据是否完整"],
])

# ===== 二、两个工作流与覆盖组件 =====
doc.add_heading("二、两个工作流与覆盖组件", level=1)
para("测试能力由两条工作流承载：图 1 描述单笔用例在平台内的执行链路（覆盖 8 类业务组件），图 2 描述测试过程的治理闭环（覆盖 4 个测试环节）。")
add_table([
    ["组件 / 环节", "图 1 执行链路", "图 2 治理闭环"],
    ["字段", "●", "—"],
    ["外部数据", "●", "—"],
    ["数据处理（ETL）", "●", "—"],
    ["函数", "●", "—"],
    ["实时指标（10 类模板，三段式验证）", "●（重点）", "—"],
    ["规则 / 规则集", "●", "—"],
    ["决策流", "●", "—"],
    ["接口服务", "●", "—"],
    ["用例设计与覆盖门禁", "—", "●"],
    ["平台执行与 Fixture 隔离", "—", "●"],
    ["报告审计（4 维度）", "—", "●"],
    ["闭环编排与最小重跑", "—", "●"],
])
para("实时指标按三段式验证：历史贡献 + 当前笔贡献 = 最终算术值；并隔离全部指标分组维度，不能只换业务 ID。", size=10)
add_fig(FIG1, "图 1  带指标的策略测试执行链路（指标值三段式验证）")
add_fig(FIG2, "图 2  天策测试工作流（执行 → 治理闭环）")

# ===== 三、测试执行 skill 与治理闭环 =====
doc.add_heading("三、测试执行 skill 与治理闭环", level=1)
para("天策策略测试由 4 个测试执行 skill 构成，遵循同一份治理合同，产出物通过稳定 caseKey 串联。")
add_table([
    ["Skill 名", "产出物", "在测试链路中的角色"],
    ["tiance-metric-testcase-generator", "testcases.json/xlsx + design-manifest.json + coverage-plan.json", "用例设计与覆盖门禁"],
    ["tiance-metric-policy-test", "executionManifest + Fixture + 15 列 Excel 报告", "平台执行与证据采集"],
    ["tiance-metric-report-checker", "auditManifest + 质量检查 Sheet", "报告审计（4 维度）"],
    ["tiance-metric-agent-loop", "convergence.json + rerun-plan.json", "治理闭环编排（8 阶段）"],
])
para("治理闭环 8 阶段串行 + 反馈回环：")
add_table([
    ["阶段", "名称", "责任 skill", "输出产物", "硬门禁"],
    ["A", "范围与备份", "agent-loop", "备份目录 + coverage-plan.json", "原 Excel 与运行归档带时间戳备份，新结果写新 run 目录"],
    ["B", "设计/覆盖门禁", "testcase-generator", "design-manifest.json + design-review.json + testcases.governed.json", "每条规则 1 正 1 反为最低基线；边界案例显式标正反极性；设计待确认不提交"],
    ["C", "平台/执行门禁", "policy-test", "platform_snapshot + precheck_result", "policyCode/policyVersion/bizType/发布状态核对，快照 ≤10 分钟；名称缺失阻断"],
    ["D", "Fixture 准备", "policy-test", "fixtureManifest + 新 fixtureRunId", "首轮人工确认库连接/隔离键/清理范围；禁止复用上一轮隔离键"],
    ["E", "逐笔执行", "policy-test", "executionManifest + segments", "每 10 条分段；401 停；--max-retry 0；--design-manifest 必传"],
    ["F", "五态复核与证据分析", "report-checker", "governed-results.json + analysis.json", "只使用通过/失败/执行阻塞/编排阻塞/无效用例；未提交计划内用例不能从统计消失"],
    ["G", "报告质量审计", "report-checker", "auditManifest + 质量检查 Sheet", "4 维度（判定一致性/数据完整性/语义重复/实际结果质量）；安全违约非空即停"],
    ["H", "完成门禁与最小重跑", "agent-loop", "convergence.json + rerun-plan.json + cleanupResultPath", "validate-rerun 通过；caseKey 精确定位；未清理不得 completed"],
])

# ===== 四、正确理解的边界 =====
doc.add_heading("四、正确理解的边界", level=1)
para("测试用例生成和平台执行完成，不等于策略已经可以直接上线。正式使用前仍需：在目标平台使用真实或受控数据执行；核对实际指标、函数和最终结果；完成业务负责人和平台负责人的验收；发布后再次回查接口和策略结果。")
para("停止条件  如果渠道、机构、字段、外部服务、名单、指标样本或其他依赖尚未准备好，测试会明确列出缺口并停止，不会使用猜测数据继续执行。平台快照超 10 分钟阻断；Fixture run ID 每轮新建，禁止复用上一轮隔离键。默认 no-mock：不生成三方超时/错误码/字段类型错误用例，不可控三方计入覆盖缺口。")
para("明确不覆盖清单")
add_table([
    ["不覆盖项", "原因", "替代方案"],
    ["独立 .zb 指标回归", "无独立指标执行入口，须绑策略", "新建 tiance-metric-standalone-test skill"],
    ["指标计算 API 直测", "现有链路只走 /noah/policyTest", "摸 salaxy/indexApi 契约后新建"],
    ["指标性能/并发/TPS", "不在功能测试范围", "JMeter/K6 独立方案"],
    ["模型效果评估", "需 A/B 与样本回溯", "独立建模评估流程"],
    ["真实第三方异常（超时/网络故障/错误码）", "默认 no-mock，不生成异常用例", "专用模拟或隔离环境"],
    ["生产策略/名单/路由修改、存量数据清理", "不属于测试授权范围", "走变更管理流程"],
])

# ===== 五、汇报结论 =====
doc.add_heading("五、汇报结论", level=1)
para("业务化概括  业务人员说明在哪个渠道、针对什么对象、需要哪些数据、按照什么条件判断、最终期望返回什么结果，系统可以协助形成测试场景、准备测试数据、执行平台验证，并输出一套可解释、可追溯的测试结论。每条用例保留稳定 caseKey、设计清单、执行证据与治理状态；最小重跑基于 caseKey 精确修改。")
para("测试结果的五态分类：")
add_table([
    ["结果", "业务含义", "下一步"],
    ["通过", "目标指标、规则和策略结果均有证据且符合预期", "纳入通过统计"],
    ["失败", "实际结果与业务预期存在已证实差异", "根因分析（数据/配置/平台/机制），产出 analysis.json"],
    ["执行阻塞", "环境/接口/数据库/平台故障或 401/超时导致本条未完成", "修复环境后按 caseKey 最小重跑"],
    ["编排阻塞", "设计门禁/平台门禁/Fixture 门禁未通过，用例未提交", "补齐前置条件后进入 B/C/D 阶段"],
    ["无效用例", "设计歧义、caseKey 冲突、覆盖重复或场景不适用", "记入 excludedScenarios 并计入覆盖缺口"],
])
para("最终判断原则  只有目标指标、目标规则和最终策略结果均有证据并符合预期，测试才判定通过。")

doc.save(OUT)
print("saved:", OUT)
