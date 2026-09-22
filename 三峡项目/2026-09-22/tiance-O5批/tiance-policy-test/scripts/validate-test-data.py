#!/usr/bin/env python3

import argparse
import copy
import json
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Error: openpyxl is required. Install it with: python3 -m pip install openpyxl", file=sys.stderr)
    sys.exit(1)


PROFILE_PATH = Path(__file__).resolve().parents[1] / "references" / "hlb_acquiring_profile.json"


def _load_profile(path=PROFILE_PATH):
    """Derive ENDPOINTS/ENUMS from the machine-readable HLB profile JSON.

    Single source of truth (O5): references/hlb_acquiring_profile.json. Shapes
    are rebuilt to match the pre-refactor literals byte-for-byte — required
    groups and enum alias/value tuples stay ordered as authored in the JSON.
    """
    if not Path(path).exists():
        raise FileNotFoundError(f"HLB profile not found at {path}")
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    endpoints = {
        name: {"path": spec["path"], "required": [tuple(group) for group in spec["required"]]}
        for name, spec in data["endpoints"].items()
    }
    enums = tuple(
        (tuple(entry["aliases"]), tuple(entry["valid"])) for entry in data["enums"]
    )
    return endpoints, enums


ENDPOINTS, ENUMS = _load_profile()


def present(value):
    return value is not None and value != ""


def map_deep(value, mapper):
    if isinstance(value, str):
        return mapper(value)
    if isinstance(value, list):
        return [map_deep(item, mapper) for item in value]
    if isinstance(value, dict):
        return {key: map_deep(item, mapper) for key, item in value.items()}
    return value


def replace_sequence(value, index):
    def replace(text):
        def repl(match):
            width = max(len(match.group(1)), len(match.group(2)))
            return str(index).zfill(width)
        return re.sub(r"\{序号(\d+)-(\d+)\}", repl, text)
    return map_deep(value, replace)


def ensure_identifiers(params, index, count):
    if count <= 1 or index <= 1:
        return params
    width = len(str(count))
    for field in ("bizid", "S_S_BIZID", "externaltransactionid"):
        if present(params.get(field)):
            params[field] = f"{params[field]}_{index:0{width}d}"
    return params


def apply_distribution(params, distribution, index):
    result = copy.deepcopy(params)
    groups_by_field = {}
    for group in distribution if isinstance(distribution, list) else []:
        if not isinstance(group, dict) or not isinstance(group.get("字段"), dict):
            continue
        try:
            count = int(group.get("数量", 0))
        except (TypeError, ValueError):
            continue
        if count < 0:
            continue
        for field, value in group["字段"].items():
            groups_by_field.setdefault(field, []).append((count, value))
    for field, groups in groups_by_field.items():
        upper = 0
        for count, value in groups:
            upper += count
            if index <= upper:
                result[field] = value
                break
    return result


def shift_date(params, index, spread_days):
    if not spread_days or index == 1 or not params.get("biztime"):
        return params
    match = re.match(r"^(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})", str(params["biztime"]))
    if not match:
        return params
    value = datetime.strptime(f"{match.group(1)} {match.group(2)}", "%Y-%m-%d %H:%M:%S")
    value -= timedelta(days=index - 1)
    params["biztime"] = value.strftime("%Y-%m-%d %H:%M:%S")
    if "transactiondate" in params:
        params["transactiondate"] = value.strftime("%Y-%m-%d")
    if "transactiontime" in params:
        params["transactiontime"] = value.strftime("%H:%M:%S")
    return params


def spread_merchant(params, item, index, row_num):
    try:
        spread = int(item.get("分散商户", 0))
        start = int(item.get("商户起始", 1))
    except (TypeError, ValueError):
        return params
    if spread <= 0:
        return params
    merchant = f"M{row_num:06d}_B{start + ((index - 1) % spread):03d}"
    if "merchantid" in params:
        params["merchantid"] = merchant
    if "C_S_MERCHANTID" in params:
        params["C_S_MERCHANTID"] = merchant
    if item.get("账号用商户"):
        if "customeracctnumber" in params:
            params["customeracctnumber"] = merchant
        if "C_S_CUSTOMERACCTNUMBER" in params:
            params["C_S_CUSTOMERACCTNUMBER"] = merchant
    return params


def expand_history(items, row_num, errors):
    expanded = []
    for item_index, item in enumerate(items, start=1):
        label = f"historical transaction {item_index}"
        if not isinstance(item, dict):
            errors.append(f"{label} must be an object")
            continue
        if isinstance(item.get("请求参数"), dict):
            expanded.append({
                "endpoint": item.get("接口地址"),
                "params": copy.deepcopy(item["请求参数"]),
                "source": label,
            })
            continue
        template = item.get("请求参数模板")
        try:
            count = int(item.get("批量生成数量", 0))
        except (TypeError, ValueError):
            count = 0
        if not isinstance(template, dict):
            errors.append(f"Missing 请求参数 or 请求参数模板 in {label}")
            continue
        if count <= 0:
            errors.append(f"批量生成数量 must be a positive integer in {label}")
            continue
        distribution = item.get("字段分布", [])
        total = sum(int(group.get("数量", 0) or 0) for group in distribution if isinstance(group, dict))
        if total > count:
            errors.append(f"字段分布数量 {total} exceeds 批量生成数量 {count} in {label}")
        for index in range(1, count + 1):
            params = replace_sequence(template, index)
            if params == template:
                params = ensure_identifiers(params, index, count)
            params = apply_distribution(params, distribution, index)
            params = shift_date(params, index, item.get("分散天数", 0))
            params = spread_merchant(params, item, index, row_num)
            expanded.append({
                "endpoint": item.get("接口地址"),
                "params": params,
                "source": f"{label} expanded item {index}",
            })
    return expanded


def endpoint_type(endpoint):
    if not endpoint:
        return None
    return next((name for name, spec in ENDPOINTS.items() if spec["path"] in str(endpoint)), None)


def validate_endpoint(endpoint, label, errors):
    if not endpoint:
        errors.append(f"Missing interface address for {label}")
        return None
    kind = endpoint_type(endpoint)
    if not kind:
        errors.append(f"Unsupported interface address for {label}: {endpoint}")
    return kind


def validate_params(params, kind, label, errors):
    if not isinstance(params, dict):
        errors.append(f"Missing request parameters in {label}")
        return
    if kind:
        for aliases in ENDPOINTS[kind]["required"]:
            if not any(present(params.get(field)) for field in aliases):
                errors.append(f"Missing required field ({' or '.join(aliases)}) in {label}")
    for aliases, values in ENUMS:
        for field in aliases:
            if present(params.get(field)) and str(params[field]) not in values:
                errors.append(f"Invalid enum value {params[field]} for {field} in {label}. Valid: {', '.join(values)}")
    for field, value in params.items():
        if not present(value):
            errors.append(f"Empty field {field} in {label}; delete empty fields")


def validate_metadata(test_data, current, execution_ready, errors, warnings):
    metadata = test_data.get("执行元数据")
    if not isinstance(metadata, dict):
        message = "Missing 执行元数据; target-rule classification is not execution-ready"
        (errors if execution_ready else warnings).append(message)
        return
    for field in ("目标规则", "目标规则集", "策略编码", "案例类型"):
        if not present(metadata.get(field)):
            errors.append(f"Missing 执行元数据.{field}")
    if present(metadata.get("案例类型")) and metadata["案例类型"] not in ("正案例", "反案例"):
        errors.append(f"Invalid 执行元数据.案例类型 {metadata['案例类型']}; valid: 正案例, 反案例")
    inclusion = metadata.get("当前笔计入方式")
    if present(inclusion) and inclusion not in ("included", "excluded", "unknown"):
        errors.append(f"Invalid 执行元数据.当前笔计入方式 {inclusion}")
    expected = metadata.get("预期当前笔字段", {})
    if not isinstance(expected, dict):
        errors.append("执行元数据.预期当前笔字段 must be an object")
        return
    for field, expected_value in expected.items():
        if field not in current:
            errors.append(f"Expected current field {field} is missing from normalized current request")
        elif current[field] != expected_value:
            errors.append(f"Current field mismatch {field}: expected {json.dumps(expected_value, ensure_ascii=False)}, actual {json.dumps(current[field], ensure_ascii=False)}")


def validate_unique_ids(current, history, errors):
    seen = set()
    transactions = [(current, "current transaction")] + [(item["params"], item["source"]) for item in history]
    for params, source in transactions:
        bizid = params.get("bizid", params.get("S_S_BIZID")) if isinstance(params, dict) else None
        if not present(bizid):
            continue
        if str(bizid) in seen:
            errors.append(f"Duplicate bizid {bizid} in {source}")
        seen.add(str(bizid))


def validate_file(excel_path, execution_ready=False):
    workbook = load_workbook(excel_path, read_only=True, data_only=False)
    sheet = workbook["测试用例"] if "测试用例" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        print("Error: Excel file has no data rows", file=sys.stderr)
        return 1
    headers = [str(value or "").strip() for value in rows[0]]
    json_column = headers.index("测试数据") if "测试数据" in headers else 13
    valid_count = invalid_count = expanded_count = 0
    all_errors = []
    all_warnings = []
    for row_num, row in enumerate(rows[1:], start=2):
        errors = []
        warnings = []
        raw = row[json_column] if json_column < len(row) else None
        if not present(raw):
            errors.append("Missing JSON data in test data column")
        elif len(str(raw)) > 32767:
            errors.append(f"JSON length {len(str(raw))} exceeds Excel limit 32767")
        test_data = None
        if not errors:
            try:
                test_data = json.loads(raw) if isinstance(raw, str) else raw
            except Exception as error:
                errors.append(f"Invalid JSON: {error}")
        if isinstance(test_data, dict):
            current = test_data.get("当前笔", test_data.get("触发交易(当前笔)"))
            endpoint = test_data.get("当前笔接口", test_data.get("接口地址"))
            history_source = test_data.get("历史数据(构造事件历史)", [])
            if not isinstance(current, dict):
                errors.append("Missing 当前笔 or 触发交易(当前笔) in JSON")
                current = {}
            if not isinstance(history_source, list):
                errors.append("历史数据(构造事件历史) must be an array")
                history_source = []
            history = expand_history(history_source, row_num, errors)
            current_type = validate_endpoint(endpoint, "current transaction", errors)
            validate_params(current, current_type, "current transaction", errors)
            for item in history:
                kind = validate_endpoint(item["endpoint"], item["source"], errors)
                validate_params(item["params"], kind, item["source"], errors)
            validate_unique_ids(current, history, errors)
            validate_metadata(test_data, current, execution_ready, errors, warnings)
            expanded_count += len(history)
        if errors:
            invalid_count += 1
            all_errors.extend((row_num, error) for error in errors)
        else:
            valid_count += 1
        all_warnings.extend((row_num, warning) for warning in warnings)
    print(f"Validating: {excel_path}\n")
    print("Validation Report")
    print(f"Total test cases: {len(rows) - 1}")
    print(f"Valid: {valid_count}")
    print(f"Invalid: {invalid_count}")
    print(f"Expanded historical transactions: {expanded_count}")
    if all_warnings:
        print(f"\nWarnings ({len(all_warnings)}):")
        for row_num, warning in all_warnings:
            print(f"  Row {row_num}: {warning}")
    if all_errors:
        print(f"\nErrors ({len(all_errors)}):")
        for row_num, error in all_errors:
            print(f"  Row {row_num}: {error}")
        return 1
    print("\nAll test cases are valid!")
    return 0


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--execution-ready", action="store_true")
    parser.add_argument("excel_file")
    args = parser.parse_args()
    path = Path(args.excel_file)
    if not path.exists():
        print(f"Error: File not found: {path}", file=sys.stderr)
        return 1
    return validate_file(path, args.execution_ready)


if __name__ == "__main__":
    sys.exit(main())
